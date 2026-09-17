"""Ashlar-branded client PPTX generator ported from Proposal Studio."""
from __future__ import annotations

import io
import re
from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from .client_analysis import plan_display_name, premium_display, client_facing_deductible, find_plan_narrative

NAVY = RGBColor(12, 26, 42)
NAVY2 = RGBColor(18, 36, 58)
INK = RGBColor(28, 37, 55)
MUTED = RGBColor(105, 117, 137)
BORDER = RGBColor(225, 231, 239)
BG = RGBColor(246, 248, 251)
WHITE = RGBColor(255, 255, 255)
PURPLE = RGBColor(108, 76, 245)
GOLD = RGBColor(231, 173, 67)
GREEN = RGBColor(21, 150, 106)
FONT = "Aptos"


def _labels(language: str) -> dict:
    greek = str(language).lower().startswith(("gr", "el")) or "greek" in str(language).lower()
    if greek:
        return {
            "subtitle": "Ανεξάρτητη συγκριτική ανάλυση και αιτιολογημένη συμβουλευτική άποψη",
            "reviewed": "ασφαλιστικές επιλογές εξετάστηκαν",
            "context": "ΠΛΑΙΣΙΟ ΠΕΛΑΤΗ",
            "what_matters": "Τι έχει σημασία σε αυτή τη σύγκριση",
            "executive": "Συνοπτική εικόνα",
            "overview": "ΣΥΝΟΨΗ ΕΠΙΛΟΓΩΝ",
            "glance": "Τα προγράμματα με μια ματιά",
            "premium": "ΑΣΦΑΛΙΣΤΡΟ",
            "annual": "ΕΤΗΣΙΟ ΟΡΙΟ",
            "area": "ΠΕΡΙΟΧΗ",
            "excess": "ΑΠΑΛΛΑΓΗ",
            "plan_review": "ΠΑΡΟΥΣΙΑΣΗ ΠΡΟΓΡΑΜΜΑΤΟΣ",
            "strengths": "Πλεονεκτήματα",
            "consider": "Σημεία προσοχής",
            "comparison": "ΣΥΓΚΡΙΤΙΚΗ ΑΠΕΙΚΟΝΙΣΗ",
            "compare_title": "Πώς συγκρίνονται οι επιλογές",
            "benefit": "Κάλυψη / όρος",
            "what_matters_sec": "ΟΥΣΙΑΣΤΙΚΕΣ ΔΙΑΦΟΡΕΣ",
            "diff_title": "Οι διαφορές που έχουν σημασία",
            "assessment": "ΠΡΟΤΑΣΗ ASHLAR",
            "preferred": "Προτεινόμενη επιλογή",
            "alternative": "Ισχυρή εναλλακτική",
            "extras": "Επιλογή με ευρύτερα πρόσθετα",
            "budget": "Οικονομικότερη επιλογή",
            "before": "ΠΡΙΝ ΠΡΟΧΩΡΗΣΕΤΕ",
            "important": "Σημαντικές επισημάνσεις",
            "next": "Επόμενα βήματα",
            "notice": "Σημαντική σημείωση",
            "footer": "Ανεξάρτητη συγκριτική ανάλυση",
            "final": "Η τελική επιλογή υπόκειται στο underwriting της ασφαλιστικής και στα ισχύοντα συμβατικά έγγραφα.",
        }
    return {
        "subtitle": "Independent plan comparison and reasoned advisory view",
        "reviewed": "insurance options reviewed",
        "context": "CLIENT CONTEXT",
        "what_matters": "What matters in this comparison",
        "executive": "Executive summary",
        "overview": "OPTIONS OVERVIEW",
        "glance": "The plans at a glance",
        "premium": "PREMIUM",
        "annual": "ANNUAL LIMIT",
        "area": "AREA",
        "excess": "EXCESS",
        "plan_review": "PLAN REVIEW",
        "strengths": "Strengths",
        "consider": "Points to consider",
        "comparison": "SIDE-BY-SIDE COMPARISON",
        "compare_title": "How the options compare",
        "benefit": "Benefit / term",
        "what_matters_sec": "WHAT MATTERS",
        "diff_title": "The differences that matter",
        "assessment": "ASHLAR RECOMMENDATION",
        "preferred": "Recommended option",
        "alternative": "Strong alternative",
        "extras": "Broader extras option",
        "budget": "Budget option",
        "before": "BEFORE YOU PROCEED",
        "important": "Important considerations",
        "next": "Next steps",
        "notice": "Important notice",
        "footer": "Independent comparative analysis",
        "final": "The final choice remains subject to insurer underwriting and the governing policy documents.",
    }


def _rect(slide, x, y, w, h, fill, radius=False, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
    else:
        shp.line.fill.background()
    return shp


def _text(slide, text, x, y, w, h, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, fit=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(.02)
    tf.margin_right = Inches(.02)
    tf.margin_top = Inches(.01)
    tf.margin_bottom = Inches(.01)
    if fit:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text or "")
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _bullets(slide, items, x, y, w, h, size=14, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(.02)
    tf.margin_right = Inches(.02)
    tf.margin_top = Inches(.01)
    tf.margin_bottom = Inches(.01)
    for i, item in enumerate(items or []):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + str(item)
        p.level = 0
        p.space_after = Pt(6)
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
    return box


def _base(prs, section="ASHLAR ASSURANCE"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    _rect(slide, 0, 0, 13.333, .09, GOLD)
    _text(slide, section, .55, .26, 4, .3, 9, True, MUTED)
    _text(slide, "Ashlar Assurance", 10.65, .26, 2.1, .3, 9, True, MUTED, PP_ALIGN.RIGHT)
    return slide


def _footer(slide, page_note="Independent comparative analysis"):
    _text(slide, page_note, .55, 7.15, 8.5, .2, 8, False, MUTED)
    _text(slide, date.today().strftime("%d %b %Y"), 11.2, 7.15, 1.55, .2, 8, False, MUTED, PP_ALIGN.RIGHT)


def _compact(value, max_chars=90):
    s = str(value or "Not specified").replace("\n", " ").strip()
    return s if len(s) <= max_chars else s[: max_chars - 1].rstrip() + "…"


def _short_plan_label(provider, plan):
    p = str(provider or "").strip()
    pl = str(plan or "").strip()
    pk = p.casefold()
    if "cigna" in pk:
        p = "Cigna"
    elif "international medical group" in pk or pk.startswith("img"):
        p = "IMG"
    elif "now health" in pk or "starr europe" in pk:
        p = "Now Health"
    elif "bupa" in pk:
        p = "Bupa Global"
    pl = re.sub(r"\s*\([^)]*\)\s*$", "", pl).strip()
    return _compact(f"{p} · {pl}".strip(" ·"), 42)


def build_pptx_bytes(*, client_analysis: dict, results: list[dict], language: str = "English") -> bytes:
    lab = _labels(language)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    _rect(s, 0, 0, .12, 7.5, GOLD)
    _text(s, "ASHLAR ASSURANCE", .75, .55, 4, .32, 10, True, GOLD)
    _text(s, client_analysis.get("report_title") or "Insurance Comparative Analysis", .75, 1.18, 11.75, 1.72, 28, True, WHITE, fit=True)
    _text(s, client_analysis.get("client_name") or "Client", .75, 3.15, 10.5, .5, 18, True, WHITE, fit=True)
    _text(s, lab["subtitle"], .75, 3.82, 9.3, .5, 12, False, RGBColor(188, 200, 215), fit=True)
    _rect(s, .7, 5.45, 4.15, .65, NAVY2, True, RGBColor(52, 74, 98))
    _text(s, f"{len(results)} {lab['reviewed']}", .95, 5.61, 3.65, .28, 12, True, WHITE)
    _text(s, date.today().strftime("%d %B %Y"), .7, 6.68, 2.6, .25, 9, False, RGBColor(160, 175, 193))

    s = _base(prs, lab["context"])
    _text(s, lab["what_matters"], .55, .8, 12, .55, 25, True)
    _text(s, client_analysis.get("client_needs_summary") or client_analysis.get("client_priorities") or "No specific priorities supplied.", .55, 1.48, 5.9, 2.2, 15, False, INK)
    _rect(s, 6.75, 1.35, 5.95, 4.85, WHITE, True, BORDER)
    _text(s, lab["executive"], 7.05, 1.68, 5.25, .45, 17, True, PURPLE)
    _text(s, client_analysis.get("executive_summary") or "", 7.05, 2.28, 5.25, 3.55, 14, False, INK)
    _footer(s)

    for start in range(0, len(results), 4):
        subset = results[start:start + 4]
        s = _base(prs, lab["overview"])
        _text(s, lab["glance"], .55, .78, 12, .55, 25, True)
        n = len(subset)
        gap = .22
        left = .55
        total_w = 12.23
        card_w = (total_w - gap * (n - 1)) / n
        for idx, result in enumerate(subset):
            a = result.get("analysis") or {}
            x = left + idx * (card_w + gap)
            _rect(s, x, 1.55, card_w, 4.92, WHITE, True, BORDER)
            _rect(s, x, 1.55, card_w, .08, PURPLE if idx % 2 == 0 else GOLD)
            _text(s, a.get("provider") or result.get("provider"), x + .22, 1.82, card_w - .44, .3, 10, True, MUTED)
            _text(s, a.get("plan_name") or result.get("target_plan") or "Plan", x + .22, 2.17, card_w - .44, .62, 18, True, INK)
            _text(s, lab["premium"], x + .22, 2.95, card_w - .44, .22, 8, True, MUTED)
            _text(s, _compact(premium_display(a), 38), x + .22, 3.20, card_w - .44, .48, 14, True)
            _text(s, lab["annual"], x + .22, 3.82, card_w - .44, .22, 8, True, MUTED)
            _text(s, _compact(a.get("annual_limit"), 38), x + .22, 4.06, card_w - .44, .55, 13, True)
            _text(s, lab["area"], x + .22, 4.73, card_w - .44, .22, 8, True, MUTED)
            _text(s, _compact(a.get("area_of_cover"), 50), x + .22, 4.98, card_w - .44, .65, 12, True)
            _text(s, lab["excess"], x + .22, 5.78, card_w - .44, .22, 8, True, MUTED)
            _text(s, _compact(client_facing_deductible(a.get("deductible_or_excess")), 48), x + .22, 6.02, card_w - .44, .34, 10, False)
        _footer(s)

    narrative_plans = client_analysis.get("plans", [])
    for result in results:
        a = result.get("analysis") or {}
        n = find_plan_narrative(narrative_plans, result)
        s = _base(prs, lab["plan_review"])
        _text(s, a.get("provider") or result.get("provider"), .55, .75, 4, .35, 11, True, PURPLE)
        _text(s, a.get("plan_name") or result.get("target_plan") or "Plan", .55, 1.12, 8.5, .7, 27, True)
        _text(s, n.get("positioning") or "", .55, 1.82, 11.7, .55, 13, False, MUTED)
        facts = [
            ("Premium", premium_display(a)),
            ("Annual limit", a.get("annual_limit")),
            ("Excess", client_facing_deductible(a.get("deductible_or_excess"))),
            ("Area", a.get("area_of_cover")),
        ]
        for i, (fact_label, val) in enumerate(facts):
            x = .55 + i * 3.05
            _rect(s, x, 2.48, 2.82, 1.1, WHITE, True, BORDER)
            _text(s, fact_label.upper(), x + .18, 2.68, 2.45, .2, 8, True, MUTED)
            _text(s, _compact(val, 44), x + .18, 2.96, 2.45, .45, 11, True)
        _text(s, n.get("summary") or "", .55, 3.95, 5.9, 1.75, 13, False)
        _text(s, lab["strengths"], 6.82, 3.95, 2.3, .35, 15, True, GREEN)
        _bullets(s, (n.get("strengths") or [])[:4], 6.82, 4.38, 5.75, 1.15, 11)
        _text(s, lab["consider"], 6.82, 5.58, 2.7, .35, 15, True, GOLD)
        _bullets(s, (n.get("considerations") or [])[:3], 6.82, 5.98, 5.75, .82, 10)
        _footer(s)

    matrix = client_analysis.get("comparison_matrix") or []
    names = [plan_display_name(r) for r in results]
    for start in range(0, len(matrix), 7):
        rows = matrix[start:start + 7]
        s = _base(prs, lab["comparison"])
        _text(s, lab["compare_title"], .55, .76, 12, .55, 25, True)
        left = .45
        top = 1.55
        first_w = 2.15
        rem = 12.45 - first_w
        col_w = rem / max(len(names), 1)
        row_h = .69
        _rect(s, left, top, first_w, row_h, NAVY)
        _text(s, lab["benefit"], left + .12, top + .18, first_w - .24, .25, 9, True, WHITE)
        for j, name in enumerate(names):
            x = left + first_w + j * col_w
            _rect(s, x, top, col_w, row_h, NAVY2)
            _text(s, _compact(name, 32), x + .09, top + .12, col_w - .18, .45, 8, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        for i, row in enumerate(rows):
            y = top + row_h * (i + 1)
            fill = WHITE if i % 2 == 0 else RGBColor(249, 250, 252)
            _rect(s, left, y, first_w, row_h, fill, False, BORDER)
            _text(s, row.get("topic"), left + .12, y + .12, first_w - .24, .45, 9, True)
            for j, name in enumerate(names):
                x = left + first_w + j * col_w
                _rect(s, x, y, col_w, row_h, fill, False, BORDER)
                _text(s, _compact((row.get("values") or {}).get(name), 54), x + .09, y + .08, col_w - .18, .52, 8, False)
        _footer(s, lab["footer"])

    diffs = client_analysis.get("key_differences") or []
    if diffs:
        for start in range(0, min(len(diffs), 6), 3):
            subset = diffs[start:start + 3]
            s = _base(prs, lab["what_matters_sec"])
            _text(s, lab["diff_title"], .62, .82, 11.8, .5, 23, True)
            y = 1.58
            for offset, d in enumerate(subset, start + 1):
                _rect(s, .62, y, 12.0, 1.48, WHITE, True, BORDER)
                _text(s, f"{offset:02d}", .86, y + .30, .42, .26, 10, True, PURPLE)
                _text(s, d.get("title") or "Difference", 1.38, y + .20, 3.0, .62, 12.5, True, fit=True)
                _text(s, d.get("analysis") or "", 4.55, y + .18, 5.0, 1.02, 9.2, False, fit=True)
                _text(s, d.get("client_impact") or "", 9.78, y + .18, 2.52, 1.02, 8.5, True, MUTED, fit=True)
                y += 1.68
            _footer(s)

    ass = client_analysis.get("ashlar_assessment") or {}
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    _text(s, lab["assessment"], .72, .48, 4.5, .28, 9.5, True, GOLD)
    _text(s, ass.get("headline") or "Our recommendation", .72, .92, 11.75, 1.02, 22.5, True, WHITE, fit=True)
    if ass.get("recommended_provider") or ass.get("recommended_plan"):
        _rect(s, .72, 2.12, 5.05, .76, PURPLE, True)
        _text(s, f"{lab['preferred']}: {_short_plan_label(ass.get('recommended_provider'), ass.get('recommended_plan'))}", .96, 2.34, 4.55, .30, 11.8, True, WHITE, fit=True)
    _bullets(s, (ass.get("reasoning") or [])[:4], .72, 3.08, 7.0, 2.72, 11.2, WHITE)

    cards = []
    if ass.get("alternative_provider") or ass.get("alternative_plan"):
        cards.append((lab["alternative"], _short_plan_label(ass.get("alternative_provider"), ass.get("alternative_plan")), ass.get("alternative_reason") or ass.get("when_the_alternative_may_be_better") or ""))
    if ass.get("extras_provider") or ass.get("extras_plan"):
        cards.append((lab["extras"], _short_plan_label(ass.get("extras_provider"), ass.get("extras_plan")), ass.get("extras_reason") or ""))
    if ass.get("budget_provider") or ass.get("budget_plan"):
        cards.append((lab["budget"], _short_plan_label(ass.get("budget_provider"), ass.get("budget_plan")), ass.get("budget_reason") or ""))
    y = 2.12
    for idx, (label, name, reason) in enumerate(cards[:3]):
        h = 1.28 if idx else 1.48
        _rect(s, 8.05, y, 4.55, h, NAVY2, True, RGBColor(53, 75, 98))
        _text(s, label, 8.35, y + .18, 3.9, .22, 9.5, True, GOLD)
        _text(s, name, 8.35, y + .45, 3.9, .34, 12.5, True, WHITE, fit=True)
        _text(s, reason, 8.35, y + .82, 3.86, h - .9, 8.2, False, RGBColor(194, 207, 221), fit=True)
        y += h + .16
    _text(s, lab["final"], .72, 6.82, 11.7, .26, 8.5, False, RGBColor(157, 174, 193), fit=True)

    s = _base(prs, lab["before"])
    _text(s, lab["important"], .55, .78, 5.9, .5, 23, True)
    _bullets(s, (client_analysis.get("important_considerations") or [])[:7], .55, 1.55, 5.8, 4.55, 12)
    _rect(s, 6.72, 1.42, 5.95, 4.9, WHITE, True, BORDER)
    _text(s, lab["next"], 7.02, 1.78, 5.2, .45, 18, True, PURPLE)
    _bullets(s, (client_analysis.get("next_steps") or [])[:6], 7.02, 2.42, 5.1, 2.55, 12)
    _text(s, lab["notice"], 7.02, 5.18, 2.3, .28, 10, True, GOLD)
    _text(s, client_analysis.get("disclaimer") or "", 7.02, 5.55, 5.1, .62, 9, False, MUTED)
    _footer(s)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
